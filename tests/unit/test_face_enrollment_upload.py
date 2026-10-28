import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import cv2
import numpy as np
from openpyxl import load_workbook

from app.faces.face_dataset import ImageRejected, discover_dataset
from app.faces.face_enrollment_upload import (
    FaceEnrollmentConflict,
    FaceEnrollmentUploadError,
    store_face_enrollment,
)


class _AcceptingAnalyzer:
    def analyze(self, _):
        return np.asarray([1.0, 0.0], dtype=np.float32), {"quality_score": 0.95}


class _RejectingAnalyzer:
    def analyze(self, _):
        raise ImageRejected("multiple_faces")


def _encoded_image() -> bytes:
    image = np.full((96, 96, 3), 120, dtype=np.uint8)
    ok, encoded = cv2.imencode(".png", image)
    if not ok:
        raise RuntimeError("could not encode test image")
    return encoded.tobytes()


class FaceEnrollmentUploadTests(unittest.TestCase):
    def test_creates_workbook_and_importable_employee_photo(self):
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.faces.face_enrollment_upload.config.FACE_DATASET_PATH",
            directory,
        ):
            result = store_face_enrollment(
                employee_code="E001",
                full_name="Example Employee",
                department="Operations",
                image_bytes=_encoded_image(),
                actor_user_id=17,
                analyzer=_AcceptingAnalyzer(),
            )

            root = Path(directory)
            image_exists = (root / result["image_path"]).is_file()
            manifest = root / "employees.xlsx"
            people = discover_dataset(root)
            workbook = load_workbook(manifest, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                headers = [cell.value for cell in next(sheet.iter_rows())]
                row = dict(zip(
                    headers, next(sheet.iter_rows(min_row=2, values_only=True)),
                ))
            finally:
                workbook.close()

        self.assertEqual(result["employee_code"], "E001")
        self.assertTrue(image_exists)
        self.assertEqual(len(people), 1)
        self.assertEqual(len(people["E001"].images), 1)
        self.assertTrue(people["E001"].images[0].consent_recorded)
        self.assertTrue(row["consent_reference"].startswith("UI-"))
        self.assertEqual(row["consent_basis"], "Consent attested during enrollment by account 17")

    def test_uploading_an_additional_photo_preserves_existing_consent(self):
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.faces.face_enrollment_upload.config.FACE_DATASET_PATH",
            directory,
        ):
            for _ in range(2):
                store_face_enrollment(
                    employee_code="E001",
                    full_name="Example Employee",
                    department="Operations",
                    image_bytes=_encoded_image(),
                    actor_user_id=17,
                    analyzer=_AcceptingAnalyzer(),
                )
            people = discover_dataset(Path(directory))

        self.assertEqual(len(people["E001"].images), 2)
        self.assertEqual(
            len({image.consent_basis for image in people["E001"].images}), 1,
        )
        self.assertEqual(
            len({image.consent_date for image in people["E001"].images}), 1,
        )

    def test_existing_employee_id_cannot_be_reassigned_to_other_details(self):
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.faces.face_enrollment_upload.config.FACE_DATASET_PATH",
            directory,
        ):
            store_face_enrollment(
                employee_code="E001",
                full_name="Example Employee",
                department="Operations",
                image_bytes=_encoded_image(),
                actor_user_id=17,
                analyzer=_AcceptingAnalyzer(),
            )
            with self.assertRaises(FaceEnrollmentConflict):
                store_face_enrollment(
                    employee_code="E001",
                    full_name="Different Employee",
                    department="Operations",
                    image_bytes=_encoded_image(),
                    actor_user_id=17,
                    analyzer=_AcceptingAnalyzer(),
                )

            people = discover_dataset(Path(directory))

        self.assertEqual(len(people["E001"].images), 1)

    def test_rejects_bad_image_and_low_quality_face_before_writing_files(self):
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.faces.face_enrollment_upload.config.FACE_DATASET_PATH",
            directory,
        ):
            with self.assertRaisesRegex(FaceEnrollmentUploadError, "readable"):
                store_face_enrollment(
                    employee_code="E001", full_name="Example Employee",
                    department="Operations", image_bytes=b"not an image",
                    actor_user_id=17, analyzer=_AcceptingAnalyzer(),
                )
            with self.assertRaisesRegex(FaceEnrollmentUploadError, "multiple faces"):
                store_face_enrollment(
                    employee_code="E001", full_name="Example Employee",
                    department="Operations", image_bytes=_encoded_image(),
                    actor_user_id=17, analyzer=_RejectingAnalyzer(),
                )

            self.assertFalse((Path(directory) / "employees.xlsx").exists())
            self.assertFalse((Path(directory) / "enrolled_photos").exists())

    def test_rejects_unsafe_employee_ids_and_spreadsheet_formulas(self):
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.faces.face_enrollment_upload.config.FACE_DATASET_PATH",
            directory,
        ):
            for employee_code, full_name in (
                ("../E001", "Example Employee"),
                ("E001", "=HYPERLINK()"),
            ):
                with self.subTest(employee_code=employee_code):
                    with self.assertRaises(FaceEnrollmentUploadError):
                        store_face_enrollment(
                            employee_code=employee_code,
                            full_name=full_name,
                            department="Operations",
                            image_bytes=_encoded_image(),
                            actor_user_id=17,
                            analyzer=_AcceptingAnalyzer(),
                        )

            self.assertFalse((Path(directory) / "employees.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
