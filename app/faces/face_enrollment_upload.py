"""Store manager-submitted, consent-attested employee face enrollments."""
import os
import re
import threading
import uuid
from datetime import date
from pathlib import Path
from typing import Any

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .. import config
from .face_dataset import (
    FaceModelUnavailable,
    ImageRejected,
    create_face_analyzer,
    discover_dataset,
    encrypt_face_crop,
    _encryption_key,
)
from .face_manifest import read_face_manifest

_UPLOAD_LOCK = threading.Lock()
_MAX_UPLOAD_BYTES = 10 * 1024 * 1024
_MAX_IMAGE_PIXELS = 25_000_000
_EMPLOYEE_CODE_PATTERN = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}\Z")
_MANIFEST_HEADERS = (
    "employee_code",
    "full_name",
    "department",
    "designation",
    "image_path",
    "consent_recorded",
    "consent_reference",
    "consent_date",
    "consent_basis",
)


class FaceEnrollmentUploadError(ValueError):
    """A submitted enrollment cannot be safely added to the roster."""


class FaceEnrollmentConflict(FaceEnrollmentUploadError):
    """The employee ID already has conflicting identity details."""


def store_face_enrollment(
        *, employee_code: str, full_name: str, department: str,
        designation: str = "",
        image_bytes: bytes, actor_user_id: int,
        analyzer: Any | None = None) -> dict[str, str]:
    employee_code = employee_code.strip()
    full_name = full_name.strip()
    department = department.strip()
    designation = designation.strip()
    if not _EMPLOYEE_CODE_PATTERN.fullmatch(employee_code):
        raise FaceEnrollmentUploadError(
            "Employee ID must use letters, numbers, dots, underscores, or hyphens"
        )
    if not full_name or len(full_name) > 160:
        raise FaceEnrollmentUploadError("Enter a name of at most 160 characters")
    if not department or len(department) > 120:
        raise FaceEnrollmentUploadError(
            "Enter a department of at most 120 characters"
        )
    if (
        full_name[0] in "=+-@" or department[0] in "=+-@"
        or (designation and designation[0] in "=+-@")
    ):
        raise FaceEnrollmentUploadError(
            "Name and department cannot start with a spreadsheet formula character"
        )
    if not image_bytes or len(image_bytes) > _MAX_UPLOAD_BYTES:
        raise FaceEnrollmentUploadError("Image must be no larger than 10 MB")

    image = cv2.imdecode(
        np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR,
    )
    if image is None or image.size == 0:
        raise FaceEnrollmentUploadError("Upload a readable JPG or PNG image")
    height, width = image.shape[:2]
    if width * height > _MAX_IMAGE_PIXELS:
        raise FaceEnrollmentUploadError("Image dimensions are too large")

    face_analyzer = analyzer or create_face_analyzer()
    try:
        face_analyzer.analyze(image)
    except ImageRejected as exc:
        raise FaceEnrollmentUploadError(
            f"Image was not accepted ({exc.reason.replace('_', ' ')})"
        ) from exc
    except FaceModelUnavailable:
        raise

    ok, encoded = cv2.imencode(
        ".jpg", image, [cv2.IMWRITE_JPEG_QUALITY, 95],
    )
    if not ok:
        raise FaceEnrollmentUploadError("Could not normalize the uploaded image")

    dataset_root = Path(config.FACE_DATASET_PATH).expanduser().resolve()
    manifest_path = dataset_root / "employees.xlsx"
    with _UPLOAD_LOCK:
        dataset_root.mkdir(parents=True, exist_ok=True)
        workbook, sheet, header_columns = _load_or_create_manifest(manifest_path)
        try:
            existing_consent = _validate_existing_employee(
                manifest_path, employee_code, full_name, department,
                designation,
            )
            consent = existing_consent or {
                "reference": f"UI-{uuid.uuid4().hex}",
                "date": date.today().isoformat(),
                "basis": (
                    "Consent attested during enrollment by account "
                    f"{actor_user_id}"
                ),
            }
            relative_image_path = Path(
                "enrolled_photos", employee_code, f"{uuid.uuid4().hex}.jpg",
            )
            image_path = dataset_root / relative_image_path
            image_path.parent.mkdir(parents=True, exist_ok=True)
            temporary_workbook = dataset_root / (
                f".employees-{uuid.uuid4().hex}.tmp.xlsx"
            )
            try:
                image_path.write_bytes(
                    encrypt_face_crop(encoded.tobytes(), _encryption_key()),
                )
                values = {
                    "employee_code": employee_code,
                    "full_name": full_name,
                    "department": department,
                    "designation": designation,
                    "image_path": relative_image_path.as_posix(),
                    "consent_recorded": True,
                    "consent_reference": consent["reference"],
                    "consent_date": consent["date"],
                    "consent_basis": consent["basis"],
                }
                next_row = sheet.max_row + 1
                for name, column in header_columns.items():
                    if name in values:
                        sheet.cell(next_row, column, values[name])
                workbook.save(temporary_workbook)
                workbook.close()
                os.replace(temporary_workbook, manifest_path)
            except Exception:
                workbook.close()
                image_path.unlink(missing_ok=True)
                temporary_workbook.unlink(missing_ok=True)
                raise
        except Exception:
            workbook.close()
            raise

    return {
        "employee_code": employee_code,
        "image_path": relative_image_path.as_posix(),
        "manifest_path": str(manifest_path),
    }


def _load_or_create_manifest(path: Path):
    if path.exists():
        try:
            workbook = load_workbook(path)
        except (InvalidFileException, OSError, RuntimeError, ValueError) as exc:
            raise FaceEnrollmentUploadError(
                "The employee workbook could not be read"
            ) from exc
        sheet = workbook.active
        names = [
            str(cell.value or "").strip().lower()
            for cell in sheet[1]
        ]
        populated_names = [name for name in names if name]
        if not populated_names or len(populated_names) != len(set(populated_names)):
            workbook.close()
            raise FaceEnrollmentUploadError(
                "The employee workbook has missing or duplicate column names"
            )
        columns = {
            name: index for index, name in enumerate(names, start=1) if name
        }
        required = {
            "employee_code", "full_name", "department", "image_path",
            "consent_recorded", "consent_reference",
        }
        if not required.issubset(columns):
            workbook.close()
            raise FaceEnrollmentUploadError(
                "The employee workbook is missing required enrollment columns"
            )
        for optional in ("designation", "consent_date", "consent_basis"):
            if optional not in columns:
                column = sheet.max_column + 1
                sheet.cell(1, column, optional)
                columns[optional] = column
        return workbook, sheet, columns

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(list(_MANIFEST_HEADERS))
    return workbook, sheet, {
        name: index for index, name in enumerate(_MANIFEST_HEADERS, start=1)
    }


def _validate_existing_employee(
        path: Path, employee_code: str, full_name: str,
        department: str, designation: str = "") -> dict[str, str] | None:
    if not path.exists():
        return None
    try:
        people = discover_dataset(path)
    except (OSError, RuntimeError, ValueError) as exc:
        raise FaceEnrollmentUploadError(
            "The existing employee workbook has invalid enrollment records"
        ) from exc
    existing = people.get(employee_code)
    if existing is None:
        return None
    if (
        existing.full_name != full_name
        or (existing.department or "") != department
        or (existing.designation or "") != designation
    ):
        raise FaceEnrollmentConflict(
            "That employee ID already exists with different details"
        )
    rows = read_face_manifest(path)[1]
    for row in rows:
        if row.get("employee_code", "").strip() != employee_code:
            continue
        if (
            row.get("consent_recorded", "").strip().lower()
            not in {"1", "true", "yes"}
            or not row.get("consent_reference", "").strip()
        ):
            raise FaceEnrollmentUploadError(
                "Existing employee record is missing its consent record"
            )
        consent_date = row.get("consent_date", "").strip() or date.today().isoformat()
        try:
            date.fromisoformat(consent_date)
        except ValueError as exc:
            raise FaceEnrollmentUploadError(
                "Existing employee record has an invalid consent date"
            ) from exc
        return {
            "reference": row["consent_reference"].strip(),
            "date": existing.consent_date.isoformat() if existing.consent_date else consent_date,
            "basis": existing.consent_basis,
        }
    return None
