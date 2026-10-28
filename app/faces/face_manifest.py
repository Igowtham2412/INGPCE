"""Read local CSV or Excel face-enrollment manifests."""
import csv
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def read_face_manifest(path: Path) -> tuple[set[str], list[dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                headers = {
                    (header or "").strip().lower()
                    for header in (reader.fieldnames or [])
                }
                rows = [
                    {
                        (key or "").strip().lower(): _cell_text(value)
                        for key, value in row.items()
                        if key is not None
                    }
                    for row in reader
                ]
        except (OSError, csv.Error, UnicodeError) as exc:
            raise ValueError("Face manifest could not be read") from exc
        return headers, rows

    if suffix != ".xlsx":
        raise ValueError("Face manifest must be a .csv or .xlsx file")

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise RuntimeError(
            "Excel manifest support is not installed; install the optional "
            "openpyxl dependency"
        ) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            first_row = next(values, ())
            headers_list = [_cell_text(value).lower() for value in first_row]
            populated_headers = [header for header in headers_list if header]
            if len(populated_headers) != len(set(populated_headers)):
                raise ValueError("Excel face manifest has duplicate column names")
            headers = {header for header in headers_list if header}
            rows = []
            for values_row in values:
                row = {
                    header: _cell_text(value)
                    for header, value in zip(headers_list, values_row)
                    if header
                }
                if any(row.values()):
                    rows.append(row)
        finally:
            workbook.close()
    except (
        OSError, BadZipFile, InvalidFileException, ValueError, KeyError,
    ) as exc:
        raise ValueError("Excel face manifest could not be read") from exc
    return headers, rows
